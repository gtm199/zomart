"""Bank statement reading, matching to shops, and automatic renewal. Doubtful payments are never renewed."""
import csv
import hashlib
import io
import re
from datetime import datetime
from decimal import Decimal, InvalidOperation

from django.db import transaction
from django.utils import timezone

from apps.shops.models import Shop
from apps.subscriptions.models import BankCredit, SubscriptionPlan
from services import subscription_service
from services.exceptions import BusinessRuleError
from services.money import money

FIELDS = ["date", "narration", "credit", "reference"]
HEADER_WORDS = ("date", "narration", "description", "particulars", "remarks", "credit", "deposit",
                "debit", "withdrawal", "balance", "ref", "utr")
DATE_FORMATS = ["%d/%m/%Y", "%d/%m/%y", "%d-%m-%Y", "%d-%m-%y", "%d-%b-%Y", "%d %b %Y", "%d-%b-%y", "%Y-%m-%d", "%d.%m.%Y"]


# ---------------- reading the file ----------------
def read_rows(uploaded):
    name = uploaded.name.lower()
    if name.endswith(".xlsx"):
        from openpyxl import load_workbook

        ws = load_workbook(uploaded, read_only=True, data_only=True).active

        def cell(c):
            if c is None:
                return ""
            if isinstance(c, float) and c.is_integer():
                c = int(c)
            return str(c).strip()

        return [[cell(c) for c in row] for row in ws.iter_rows(values_only=True)]
    raw = uploaded.read()
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            text = raw.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    return [[c.strip() for c in row] for row in csv.reader(io.StringIO(text))]


def find_header(rows):
    """Bank files often start with account details; the header is the first row that looks like column names."""
    for i, row in enumerate(rows[:40]):
        if sum(any(w in c.lower() for w in HEADER_WORDS) for c in row) >= 3:
            return i
    return 0


def _guess(headers, words, avoid=()):
    for i, h in enumerate(headers):
        low = h.lower()
        if any(w in low for w in words) and not any(a in low for a in avoid):
            return i
    return None


def guess_columns(headers):
    date = _guess(headers, ["date"], avoid=["value"])
    if date is None:
        date = _guess(headers, ["date"])
    return {
        "date": date,
        "narration": _guess(headers, ["narration", "description", "particulars", "remarks", "details"]),
        "credit": _guess(headers, ["credit", "deposit"], avoid=["card"]),
        "reference": _guess(headers, ["utr", "ref"]),
    }


def _amount(text):
    cleaned = re.sub(r"[^\d.\-]", "", text or "")
    if not cleaned:
        return None
    try:
        return money(Decimal(cleaned))
    except InvalidOperation:
        return None


def _date(text):
    for fmt in DATE_FORMATS:
        try:
            return datetime.strptime((text or "").strip(), fmt).date()
        except ValueError:
            continue
    return None


def parse_credits(rows, header_index, mapping):
    """Only money received (credits). Debits and empty rows are skipped."""
    out = []
    for row in rows[header_index + 1:]:
        def cell(key):
            i = mapping.get(key)
            return row[i].strip() if i is not None and i < len(row) else ""

        amount = _amount(cell("credit"))
        if not amount or amount <= 0:
            continue
        narration = cell("narration")
        reference = cell("reference")
        if not reference or set(reference) <= {"0"}:
            m = re.search(r"\b\d{12}\b", narration)  # a UPI reference inside the narration
            reference = m.group(0) if m else ""
        if not reference:
            reference = "H" + hashlib.sha1(f"{cell('date')}|{amount}|{narration}".encode()).hexdigest()[:20]
        out.append({"reference": reference[:80], "date": _date(cell("date")), "amount": amount, "narration": narration[:500]})
    return out


# ---------------- matching ----------------
def _normalize(text):
    return re.sub(r"[^A-Z0-9]", "", (text or "").upper())


def _options(amount, shop, plans):
    return [(p, n) for p in plans for n in range(1, 13)
            if subscription_service.amount_for(p, n, shop) == amount]


def match(amount, narration, shops, plans):
    """Returns (shop, (plan, months) or None, how, note)."""
    text = _normalize(narration)
    by_code = [s for s in shops if s.payment_code and _normalize(s.payment_code) in text]
    if len(by_code) == 1:
        shop = by_code[0]
        opts = _options(amount, shop, plans)
        if len(opts) == 1:
            return shop, opts[0], "code", ""
        return shop, None, "code", f"Shop code found, but {amount} is not one of this shop's exact amounts"
    if len(by_code) > 1:
        return None, None, "", "More than one shop code in the narration"
    found = [(s, o) for s in shops if s.pay_paise for o in _options(amount, s, plans)]
    if len(found) == 1:
        return found[0][0], found[0][1], "amount", ""
    if len(found) > 1:
        return None, None, "", "Amount fits several shops: " + ", ".join(sorted({s.name for s, _ in found}))
    return None, None, "", "No shop code and the amount matches no shop"


def process(credits, user, apply=False):
    """Preview (apply=False) or save and renew (apply=True). A reference number can only ever be used once."""
    shops = [subscription_service.ensure_payment_identity(s) for s in Shop.objects.filter(is_active=True)]
    plans = list(SubscriptionPlan.objects.filter(is_active=True, price_monthly__gt=0))
    already = set(BankCredit.objects.filter(reference__in=[c["reference"] for c in credits]).values_list("reference", flat=True))
    seen, results = set(), []
    for c in credits:
        r = dict(c, shop=None, plan=None, months=0, how="", note="")
        if c["reference"] in already:
            r["outcome"] = "seen"
        elif c["reference"] in seen:
            r["outcome"] = "duplicate_in_file"
        else:
            shop, option, how, note = match(c["amount"], c["narration"], shops, plans)
            r.update(shop=shop, how=how, note=note)
            if shop and option:
                r.update(plan=option[0], months=option[1], outcome="renew")
            else:
                r["outcome"] = "review"
        seen.add(c["reference"])
        results.append(r)
    if apply:
        with transaction.atomic():
            for r in results:
                if r["outcome"] not in ("renew", "review"):
                    continue
                credit = BankCredit.objects.create(
                    reference=r["reference"], credit_date=r["date"], amount=r["amount"], narration=r["narration"],
                    shop=r["shop"], months=r["months"] or 1, matched_by=r["how"], review_note=r["note"],
                    status=BankCredit.Status.REVIEW)
                if r["outcome"] == "renew":
                    expiry = subscription_service.renew_shop(r["shop"], r["plan"], r["months"])
                    credit.status, credit.new_expiry = BankCredit.Status.RENEWED, expiry
                    credit.processed_by, credit.processed_at = user, timezone.now()
                    credit.save()
    return results


def confirm_manually(credit, user):
    """You looked at it and it is right: renew the shop named on the row for the months entered."""
    if credit.status != BankCredit.Status.REVIEW:
        raise BusinessRuleError("This payment was already handled.", "already_handled")
    if credit.shop_id is None or credit.months < 1:
        raise BusinessRuleError("Choose the shop and the number of months first.", "incomplete")
    plan = subscription_service.plan_for_manual(credit.shop)
    if plan is None:
        raise BusinessRuleError("No paid plan exists. Add one under Subscription plans.", "no_plan")
    with transaction.atomic():
        credit.new_expiry = subscription_service.renew_shop(credit.shop, plan, credit.months)
        credit.status, credit.matched_by = BankCredit.Status.RENEWED, "manual"
        credit.processed_by, credit.processed_at = user, timezone.now()
        credit.save()